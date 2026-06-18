from fpdf import FPDF
from docx import Document
from tools.write.models import ResearchReport

def generate_markdown(report: ResearchReport) -> str:
    md = f"# {report.title}\n\n"
    md += f"## Executive Summary\n{report.executive_summary.summary}\n\n"
    md += "### Key Findings\n"
    for finding in report.executive_summary.key_findings:
        md += f"- {finding}\n"
    md += "\n"

    for section in report.sections:
        md += f"## {section.title}\n{section.content}\n\n"
        if section.citations:
            md += "#### Citations\n"
            for cite in section.citations:
                md += f"- [{cite.id}] {cite.text} ({cite.url or 'N/A'})\n"
            md += "\n"

    md += "## Bibliography\n"
    for cite in report.bibliography:
        md += f"- [{cite.id}] {cite.text} - {cite.url or 'N/A'}\n"

    return md

def generate_pdf(report: ResearchReport) -> bytes:
    from fpdf.enums import XPos, YPos
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 10, report.title, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

    pdf.set_font("helvetica", "B", 14)
    pdf.cell(0, 10, "Executive Summary", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("helvetica", "", 12)
    pdf.multi_cell(0, 10, report.executive_summary.summary)

    for section in report.sections:
        pdf.set_font("helvetica", "B", 14)
        pdf.cell(0, 10, section.title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("helvetica", "", 12)
        pdf.multi_cell(0, 10, section.content)

    return bytes(pdf.output())

def generate_docx(report: ResearchReport) -> bytes:
    doc = Document()
    doc.add_heading(report.title, 0)

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(report.executive_summary.summary)

    for section in report.sections:
        doc.add_heading(section.title, level=1)
        doc.add_paragraph(section.content)

    from io import BytesIO
    target = BytesIO()
    doc.save(target)
    return target.getvalue()
